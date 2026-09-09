"""Programmatic report layout with explicitly owned pt-BR presentation text."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any, TypedDict

from reporting import locale_pt_br as pt_br

from reporting.engine import EXPECTED_RULESET


class _DesignTokens(TypedDict):
    """Describe existing color strings and numeric geometry without converting values."""

    navy: str
    teal: str
    amber: str
    cream: str
    light_gray: str
    border: str
    text: str
    slate: str
    pale_blue: str
    white: str
    a4_mm: tuple[int, int]
    cover_left_mm: float
    content_width_mm: float


DESIGN: _DesignTokens = {
    "navy": "0B1F33",
    "teal": "0F766E",
    "amber": "A16207",
    "cream": "FFFAEB",
    "light_gray": "F2F4F7",
    "border": "D0D5DD",
    "text": "17212B",
    "slate": "667085",
    "pale_blue": "D7E3EF",
    "white": "FFFFFF",
    "a4_mm": (210, 297),
    "cover_left_mm": 18.0,
    "content_width_mm": 174.0,
}


def _safe(value: Any, default: str = pt_br.UNAVAILABLE) -> str:
    """Render a display value with the existing absent-value substitution."""
    if value is None or value == "":
        return default
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value)


def _register_pdf_fonts() -> tuple[str, str, str]:
    """Use the same DejaVu/Carlito families found in the v3.0 reference PDFs when present.

    PDF embeds the selected TrueType fonts. A deterministic Helvetica fallback keeps the
    renderer operational in minimal environments; the artifact metadata exposes the
    actual family so visual QA can detect fallback rather than silently claiming parity.
    """
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    candidates = {
        "regular": [
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/share/fonts/dejavu/DejaVuSans.ttf",
        ],
        "bold": [
            "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
            "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf",
        ],
        "section": [
            "/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf",
            "/usr/share/fonts/truetype/carlito/Carlito-Bold.ttf",
        ],
    }
    found = {
        key: next((path for path in paths if Path(path).is_file()), None)
        for key, paths in candidates.items()
    }
    try:
        if found["regular"] and found["bold"]:
            pdfmetrics.registerFont(TTFont("GenomaSans", found["regular"]))
            pdfmetrics.registerFont(TTFont("GenomaSansBold", found["bold"]))
            if found["section"]:
                pdfmetrics.registerFont(TTFont("GenomaSectionBold", found["section"]))
                return "GenomaSans", "GenomaSansBold", "GenomaSectionBold"
            return "GenomaSans", "GenomaSansBold", "GenomaSansBold"
    except Exception:
        pass
    return "Helvetica", "Helvetica-Bold", "Helvetica-Bold"


def _hex(value: str):
    """Resolve an existing hexadecimal design token as a ReportLab color."""
    from reportlab.lib.colors import HexColor

    return HexColor("#" + value)


def _report_sections(rendered: dict[str, Any]) -> list[str]:
    """Extract report-specific headings without changing their order."""
    reserved = {
        pt_br.PURPOSE,
        pt_br.AUDIENCE,
        pt_br.SUMMARY,
        pt_br.FINDINGS,
        pt_br.EXECUTION_MANIFEST,
        pt_br.SOURCES,
        pt_br.LIMITATIONS,
    }
    result: list[str] = []
    for line in rendered.get("markdown", "").splitlines():
        if line.startswith("## "):
            name = line[3:].strip()
            if name not in reserved and name not in result:
                result.append(name)
    return result


def _pdf(rendered: dict[str, Any], path: Path) -> dict[str, Any]:
    """Write the existing programmatic PDF layout with pt-BR presentation text."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        KeepInFrame,
        PageBreak,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    metadata = rendered["metadata"]
    data = rendered.get("data", {})
    code = metadata["code"]
    accent = metadata.get("accent") or DESIGN["teal"]
    regular, bold, section_bold = _register_pdf_fonts()

    doc = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title=metadata["title"],
        author="GENOMA",
        subject="GENOMA v3.0 deterministic genomic report",
    )
    styles = getSampleStyleSheet()
    styles.add(
        ParagraphStyle(
            name="TopCode",
            fontName=bold,
            fontSize=7,
            leading=8.5,
            textColor=_hex(accent),
            spaceAfter=0,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Suite",
            fontName=bold,
            fontSize=8.5,
            leading=10.5,
            textColor=_hex(accent),
            spaceAfter=0,
        )
    )
    styles.add(
        ParagraphStyle(
            name="HeroTitle",
            fontName=bold,
            fontSize=27,
            leading=31.5,
            textColor=colors.white,
            spaceAfter=5 * mm,
        )
    )
    styles.add(
        ParagraphStyle(
            name="HeroTag",
            fontName=regular,
            fontSize=12.5,
            leading=16,
            textColor=_hex(DESIGN["pale_blue"]),
        )
    )
    styles.add(
        ParagraphStyle(
            name="MetaLabel",
            fontName=bold,
            fontSize=6.5,
            leading=8,
            textColor=_hex(DESIGN["slate"]),
        )
    )
    styles.add(
        ParagraphStyle(
            name="MetaValue", fontName=bold, fontSize=8.0, leading=9.5, textColor=_hex(accent)
        )
    )
    styles.add(
        ParagraphStyle(
            name="Section",
            fontName=section_bold,
            fontSize=12,
            leading=14.5,
            textColor=_hex(accent),
            spaceBefore=4.0 * mm,
            spaceAfter=2.2 * mm,
        )
    )
    styles.add(
        ParagraphStyle(
            name="H2",
            fontName=bold,
            fontSize=12.2,
            leading=15.0,
            textColor=_hex(accent),
            spaceBefore=4.0 * mm,
            spaceAfter=2.0 * mm,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Body",
            fontName=regular,
            fontSize=9.3,
            leading=13.8,
            textColor=_hex(DESIGN["text"]),
            spaceAfter=2.5 * mm,
        )
    )
    styles.add(
        ParagraphStyle(
            name="Small",
            fontName=regular,
            fontSize=7.0,
            leading=9.5,
            textColor=_hex(DESIGN["slate"]),
        )
    )
    styles.add(
        ParagraphStyle(
            name="Callout", fontName=bold, fontSize=8.6, leading=12, textColor=_hex(DESIGN["amber"])
        )
    )
    styles.add(
        ParagraphStyle(
            name="WhiteSmall", fontName=bold, fontSize=7.5, leading=9.5, textColor=colors.white
        )
    )

    story: list[Any] = []
    story.append(Paragraph(f"GENOMA  /  {code}{pt_br.GENERATED_REPORT_SUFFIX}", styles["TopCode"]))
    story.append(Spacer(1, 3 * mm))
    story.append(
        Table(
            [[""]],
            colWidths=[174 * mm],
            rowHeights=[0.35 * mm],
            style=TableStyle([("BACKGROUND", (0, 0), (-1, -1), _hex(DESIGN["border"]))]),
        )
    )
    story.append(Spacer(1, 14 * mm))
    story.append(Paragraph(f"{pt_br.SUITE_PREFIX}{code}", styles["Suite"]))
    story.append(Spacer(1, 6 * mm))

    hero = Table(
        [
            [Paragraph(metadata["title"], styles["HeroTitle"])],
            [Paragraph(metadata["tagline"], styles["HeroTag"])],
        ],
        colWidths=[174 * mm],
    )
    hero.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), _hex(DESIGN["navy"])),
                ("LEFTPADDING", (0, 0), (-1, -1), 6 * mm),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6 * mm),
                ("TOPPADDING", (0, 0), (-1, 0), 10 * mm),
                ("BOTTOMPADDING", (0, 1), (-1, -1), 9 * mm),
            ]
        )
    )
    story.append(hero)
    story.append(Spacer(1, 8 * mm))

    meta_rows = [
        [
            Paragraph(pt_br.TYPE_LABEL, styles["MetaLabel"]),
            Paragraph(pt_br.SCOPE_LABEL, styles["MetaLabel"]),
            Paragraph(pt_br.STATUS_LABEL, styles["MetaLabel"]),
        ],
        [
            Paragraph(pt_br.GENOMIC_RESULT, styles["MetaValue"]),
            Paragraph(pt_br.INDIVIDUAL_CASE, styles["MetaValue"]),
            Paragraph(
                _safe(data.get("post_deployment_status"), pt_br.PENDING), styles["MetaValue"]
            ),
        ],
    ]
    meta = Table(meta_rows, colWidths=[58 * mm, 58 * mm, 58 * mm])
    meta.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), _hex(DESIGN["light_gray"])),
                ("LINEABOVE", (0, 0), (-1, 0), 1.2, _hex(accent)),
                ("BOX", (0, 0), (-1, -1), 0.35, _hex(DESIGN["border"])),
                ("INNERGRID", (0, 0), (-1, -1), 0.25, _hex(DESIGN["border"])),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 3.2 * mm),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3.2 * mm),
                ("TOPPADDING", (0, 0), (-1, -1), 2.8 * mm),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 2.8 * mm),
            ]
        )
    )
    story.append(meta)
    story.append(Spacer(1, 7 * mm))

    story.append(Paragraph(pt_br.PURPOSE, styles["Section"]))
    story.append(Paragraph(metadata["purpose"], styles["Body"]))
    story.append(Paragraph(pt_br.AUDIENCE, styles["Section"]))
    story.append(Paragraph(metadata["audience"], styles["Body"]))

    safety = Table(
        [[Paragraph(pt_br.CONTROLLED_PUBLICATION_NOTICE, styles["Callout"])]], colWidths=[174 * mm]
    )
    safety.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, -1), _hex(DESIGN["cream"])),
                ("BOX", (0, 0), (-1, -1), 0.6, _hex(DESIGN["amber"])),
                ("LEFTPADDING", (0, 0), (-1, -1), 4.5 * mm),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4.5 * mm),
                ("TOPPADDING", (0, 0), (-1, -1), 3.5 * mm),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3.5 * mm),
            ]
        )
    )
    story.append(Spacer(1, 5 * mm))
    story.append(safety)

    body_flow: list[Any] = []
    body_flow.append(Paragraph(pt_br.SUMMARY, styles["H2"]))
    body_flow.append(Paragraph(_safe(data.get("summary")), styles["Body"]))
    section_data = data.get("sections") if isinstance(data.get("sections"), dict) else {}
    for name in _report_sections(rendered):
        body_flow.append(Paragraph(name, styles["H2"]))
        body_flow.append(Paragraph(_safe(section_data.get(name)), styles["Body"]))

    body_flow.append(Paragraph(pt_br.FINDINGS, styles["H2"]))
    findings = data.get("findings") if isinstance(data.get("findings"), list) else []
    if not findings:
        body_flow.append(Paragraph(pt_br.UNAVAILABLE, styles["Body"]))
    else:
        for finding in findings:
            if not isinstance(finding, dict):
                continue
            rows = [
                [
                    Paragraph(pt_br.FIELD_LABEL, styles["WhiteSmall"]),
                    Paragraph(pt_br.VALUE_LABEL, styles["WhiteSmall"]),
                ]
            ]
            for key, value in finding.items():
                rows.append(
                    [
                        Paragraph(key.replace("_", " ").upper(), styles["Small"]),
                        Paragraph(_safe(value), styles["Body"]),
                    ]
                )
            table = Table(rows, colWidths=[48 * mm, 126 * mm], repeatRows=1)
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), _hex(accent)),
                        ("GRID", (0, 0), (-1, -1), 0.3, _hex(DESIGN["border"])),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 2.8 * mm),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 2.8 * mm),
                        ("TOPPADDING", (0, 0), (-1, -1), 2.2 * mm),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 2.2 * mm),
                    ]
                )
            )
            body_flow.extend([table, Spacer(1, 3 * mm)])

    for heading, value in (
        (pt_br.SOURCES, data.get("sources")),
        (pt_br.LIMITATIONS, data.get("limitations")),
        (pt_br.EXECUTION_MANIFEST, data.get("execution_manifest")),
    ):
        body_flow.append(Paragraph(heading, styles["H2"]))
        body_flow.append(
            Paragraph(
                _safe(value),
                styles["Body"] if heading != pt_br.EXECUTION_MANIFEST else styles["Small"],
            )
        )

    if metadata["report_id"] == "10":
        story.append(Spacer(1, 4 * mm))
        story.append(KeepInFrame(174 * mm, 86 * mm, body_flow, mode="shrink"))
    else:
        story.append(PageBreak())
        story.extend(body_flow)

    def page_decor(canvas, _doc):
        """Draw the existing confidential page footer without changing geometry."""
        canvas.saveState()
        canvas.setFillColor(_hex(DESIGN["slate"]))
        canvas.setFont(regular, 7)
        canvas.drawString(18 * mm, 7 * mm, f"GENOMA / {code} / v3.0")
        canvas.drawRightString(
            A4[0] - 18 * mm, 7 * mm, f"{pt_br.CONFIDENTIAL_PAGE_PREFIX}{canvas.getPageNumber()}"
        )
        canvas.restoreState()

    doc.build(story, onFirstPage=page_decor, onLaterPages=page_decor)
    return {
        "pdf_font_regular": regular,
        "pdf_font_bold": bold,
        "pdf_font_section": section_bold,
        "accent": accent,
        "a4": True,
    }


def _set_cell_shading(cell, fill: str) -> None:
    """Set the existing cell background in the editable document."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_border(cell, color: str, size: str = "3") -> None:
    """Apply the existing border parameters to every cell edge."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn("w:" + edge))
        if element is None:
            element = OxmlElement("w:" + edge)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def _set_repeat_table_header(row) -> None:
    """Keep the document table header repeated across pages."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def _docx(rendered: dict[str, Any], path: Path) -> dict[str, Any]:
    """Write the existing editable document without changing layout or field values."""
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    metadata = rendered["metadata"]
    data = rendered.get("data", {})
    accent = metadata.get("accent") or DESIGN["teal"]
    code = metadata["code"]
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(17)
    section.bottom_margin = Mm(17)
    section.left_margin = Mm(18)
    section.right_margin = Mm(18)

    normal = doc.styles["Normal"]
    normal.font.name = "DejaVu Sans"
    normal.font.size = Pt(9.3)
    normal.font.color.rgb = RGBColor.from_string(DESIGN["text"])
    for style_name, size, color in (
        ("Title", 27, DESIGN["navy"]),
        ("Heading 1", 12.2, accent),
        ("Heading 2", 11, accent),
    ):
        style = doc.styles[style_name]
        style.font.name = "DejaVu Sans"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)

    header = section.header.paragraphs[0]
    header.text = f"GENOMA  /  {code}{pt_br.GENERATED_REPORT_SUFFIX}"
    run = header.runs[0]
    run.font.name = "DejaVu Sans"
    run.font.size = Pt(7)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(accent)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(pt_br.CONFIDENTIAL_PAGE_PREFIX)
    run.font.name = "DejaVu Sans"
    run.font.size = Pt(7)
    run.font.color.rgb = RGBColor.from_string(DESIGN["slate"])
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    suite = doc.add_paragraph()
    suite.paragraph_format.space_before = Pt(18)
    suite.paragraph_format.space_after = Pt(9)
    run = suite.add_run(f"{pt_br.SUITE_PREFIX}{code}")
    run.font.name = "DejaVu Sans"
    run.font.size = Pt(8.5)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(accent)

    hero = doc.add_table(rows=2, cols=1)
    for cell in (hero.cell(0, 0), hero.cell(1, 0)):
        _set_cell_shading(cell, DESIGN["navy"])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = hero.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_before = Pt(11)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(metadata["title"])
    r.font.name = "DejaVu Sans"
    r.font.size = Pt(27)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(DESIGN["white"])
    p = hero.cell(1, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(11)
    r = p.add_run(metadata["tagline"])
    r.font.name = "DejaVu Sans"
    r.font.size = Pt(12.5)
    r.font.color.rgb = RGBColor.from_string(DESIGN["pale_blue"])

    meta = doc.add_table(rows=2, cols=3)
    labels = (pt_br.TYPE_LABEL, pt_br.SCOPE_LABEL, pt_br.STATUS_LABEL)
    values = (
        pt_br.GENOMIC_RESULT,
        pt_br.INDIVIDUAL_CASE,
        _safe(data.get("post_deployment_status"), pt_br.PENDING),
    )
    for row in meta.rows:
        for cell in row.cells:
            _set_cell_shading(cell, DESIGN["light_gray"])
            _set_cell_border(cell, DESIGN["border"])
    for index, label in enumerate(labels):
        r = meta.cell(0, index).paragraphs[0].add_run(label)
        r.font.name = "DejaVu Sans"
        r.font.size = Pt(6.5)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(DESIGN["slate"])
    for index, value in enumerate(values):
        r = meta.cell(1, index).paragraphs[0].add_run(value)
        r.font.name = "DejaVu Sans"
        r.font.size = Pt(8)
        r.font.bold = True
        r.font.color.rgb = RGBColor.from_string(accent)

    def add_heading(text: str) -> None:
        """Append a heading with the existing spacing and style."""
        p = doc.add_paragraph(text, style="Heading 1")
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)

    def add_body(value: Any) -> None:
        """Append the safely rendered body value with the existing spacing."""
        p = doc.add_paragraph(_safe(value))
        p.paragraph_format.space_after = Pt(5)

    add_heading(pt_br.PURPOSE)
    add_body(metadata["purpose"])
    add_heading(pt_br.AUDIENCE)
    add_body(metadata["audience"])

    safety = doc.add_table(rows=1, cols=1)
    _set_cell_shading(safety.cell(0, 0), DESIGN["cream"])
    _set_cell_border(safety.cell(0, 0), DESIGN["amber"], "5")
    r = safety.cell(0, 0).paragraphs[0].add_run(pt_br.CONTROLLED_PUBLICATION_NOTICE)
    r.font.name = "DejaVu Sans"
    r.font.size = Pt(8.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor.from_string(DESIGN["amber"])

    add_heading(pt_br.SUMMARY)
    add_body(data.get("summary"))
    section_data = data.get("sections") if isinstance(data.get("sections"), dict) else {}
    for name in _report_sections(rendered):
        add_heading(name)
        add_body(section_data.get(name))

    add_heading(pt_br.FINDINGS)
    findings = data.get("findings") if isinstance(data.get("findings"), list) else []
    if not findings:
        add_body(pt_br.UNAVAILABLE)
    else:
        for finding in findings:
            if not isinstance(finding, dict):
                continue
            table = doc.add_table(rows=1, cols=2)
            _set_repeat_table_header(table.rows[0])
            for idx, label in enumerate((pt_br.FIELD_LABEL, pt_br.VALUE_LABEL)):
                cell = table.rows[0].cells[idx]
                _set_cell_shading(cell, accent)
                r = cell.paragraphs[0].add_run(label)
                r.font.name = "DejaVu Sans"
                r.font.size = Pt(7.5)
                r.font.bold = True
                r.font.color.rgb = RGBColor.from_string(DESIGN["white"])
            for key, value in finding.items():
                cells = table.add_row().cells
                cells[0].text = key.replace("_", " ").upper()
                cells[1].text = _safe(value)
                for cell in cells:
                    _set_cell_border(cell, DESIGN["border"])

    for heading, value in (
        (pt_br.SOURCES, data.get("sources")),
        (pt_br.LIMITATIONS, data.get("limitations")),
        (pt_br.EXECUTION_MANIFEST, data.get("execution_manifest")),
    ):
        add_heading(heading)
        add_body(value)

    doc.core_properties.title = metadata["title"]
    doc.core_properties.subject = f"GENOMA v3.0 / {code} / deterministic genomic report"
    # Ruleset identity comes from the single canonical source; never restate it here.
    doc.core_properties.keywords = f"GENOMA,{code},{EXPECTED_RULESET['version']},genomics,report"
    doc.save(path)
    return {"docx_font": "DejaVu Sans", "accent": accent, "a4": True}


def write_editorial_bundle(
    rendered: dict[str, Any], output_dir: Path, *, stem: str | None = None
) -> dict[str, Path]:
    """Write programmatic PDF/DOCX and runtime metadata for the prepared input."""
    output_dir.mkdir(parents=True, exist_ok=True)
    metadata = rendered["metadata"]
    stem = stem or f"{metadata['report_id']}-{metadata['slug']}"
    pdf = output_dir / f"{stem}.pdf"
    docx = output_dir / f"{stem}.docx"
    pdf_runtime = _pdf(rendered, pdf)
    docx_runtime = _docx(rendered, docx)
    runtime = {
        "schema": "genoma-editorial-runtime-v1",
        "report_id": metadata["report_id"],
        "code": metadata["code"],
        "accent": metadata["accent"],
        "pdf": pdf_runtime,
        "docx": docx_runtime,
        "visual_reference": "GENOMA model suite v3.0",
    }
    (output_dir / f"{stem}.editorial.json").write_text(
        json.dumps(runtime, ensure_ascii=False, indent=2, sort_keys=True) + "\n", encoding="utf-8"
    )
    return {"pdf": pdf, "docx": docx}


__all__ = ["DESIGN", "write_editorial_bundle"]
